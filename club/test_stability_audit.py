"""Regression evidence from the 2026-10-29 stability review."""
import json
from datetime import date,timedelta
from types import SimpleNamespace
from unittest.mock import patch,Mock
from django.contrib.auth.models import User
from django.db import close_old_connections
from django.test import TestCase,SimpleTestCase,TransactionTestCase,Client
from django.utils import timezone
from . import test_community as fixtures
from .models import (CheckinRecord,StreakReward,WeeklyStreakGrant,CurrencyBalance,CurrencyTransaction,
    ShopItem,UserInventory,Purchase,Team,TeamMembership,ContestSyncState,ContestParticipation,UserProfile)
from .currency_services import settle_weekly_streak,get_current_streak,purchase_item
from .contest_rewards import reward_for,REWARD_START


class AuditWeekTests(TestCase):
    setUp=fixtures.CommunityTests.setUp
    image=fixtures.CommunityTests.image
    submit=fixtures.CommunityTests.submit

    def enable(self):
        self.team.currency_enabled=True;self.team.save()
        self.monday=date(2026,10,19)
        StreakReward.objects.create(team=self.team,days=3,reward=1)
        StreakReward.objects.create(team=self.team,days=7,reward=2)

    def checkins(self,offsets):
        for offset in offsets:CheckinRecord.objects.create(team=self.team,member=self.author_member,checkin_date=self.monday+timedelta(days=offset))

    def test_two_runs_in_one_week_each_threshold_only_once_then_configuration_change(self):
        self.enable();self.checkins([0,1,2])
        self.assertEqual(settle_weekly_streak(self.author_member,self.monday),1)
        self.checkins([4,5,6])
        self.assertEqual(settle_weekly_streak(self.author_member,self.monday),0)
        self.checkins([3])
        self.assertEqual(settle_weekly_streak(self.author_member,self.monday),2)
        self.assertEqual(settle_weekly_streak(self.author_member,self.monday),0)
        StreakReward.objects.filter(team=self.team).delete()
        StreakReward.objects.create(team=self.team,days=3,reward=10)
        self.assertEqual(settle_weekly_streak(self.author_member,self.monday),0)
        self.assertEqual(CurrencyBalance.objects.get(member=self.author_member).balance,3)
        self.assertEqual(WeeklyStreakGrant.objects.filter(member=self.author_member).count(),2)

    def test_natural_week_boundary_and_future_data_do_not_grant_early(self):
        self.enable();self.checkins([4,5,6,7,8])
        self.assertEqual(settle_weekly_streak(self.author_member,self.monday),1)
        self.assertEqual(settle_weekly_streak(self.author_member,self.monday+timedelta(days=7)),0)
        self.checkins([9])
        self.assertEqual(settle_weekly_streak(self.author_member,self.monday+timedelta(days=7)),1)
        with patch('django.utils.timezone.now',return_value=timezone.make_aware(timezone.datetime(2026,10,20,10))):
            self.assertEqual(settle_weekly_streak(self.author_member,self.monday),0)

    def test_weekly_grant_survives_member_deletion(self):
        self.enable();self.checkins([0,1,2])
        self.assertEqual(settle_weekly_streak(self.author_member,self.monday),1)
        self.author_member.delete()
        self.author_member=TeamMembership.objects.create(team=self.team,user=self.author,display_name='Rejoined')
        self.checkins([0,1,2])
        self.assertEqual(settle_weekly_streak(self.author_member,self.monday),0)
        self.assertEqual(WeeklyStreakGrant.objects.filter(team=self.team,user=self.author).count(),1)

    def test_weekly_credit_failure_rolls_back_grant(self):
        self.enable();self.checkins([0,1,2])
        with patch('club.currency_services.add_currency',side_effect=ValueError('storage limit')):
            with self.assertRaises(ValueError):settle_weekly_streak(self.author_member,self.monday)
        self.assertFalse(WeeklyStreakGrant.objects.filter(member=self.author_member).exists())
        self.assertEqual(settle_weekly_streak(self.author_member,self.monday),1)

    def test_training_submission_aligns_balance_calendar_dashboard_and_rank_streak(self):
        self.enable();self.submit()
        balance=self.client.get('/api/currency/balance').json()
        board=self.client.get('/api/dashboard').json()
        calendar=self.client.get('/api/currency/checkin/calendar').json()
        ranking=self.client.get('/api/leaderboard?scope=checkin').json()['rows']
        self.assertEqual(balance['streak'],1)
        self.assertEqual(board['streakDays'],1)
        self.assertTrue(calendar['calendar'][timezone.localdate().isoformat()]['checked_in'])
        self.assertEqual(next(r['value'] for r in ranking if r['isCurrent']),1)

    def test_daily_reward_on_first_submission_only_and_failure_is_atomic(self):
        self.enable();self.team.daily_checkin_reward=3;self.team.save()
        self.submit();self.submit()
        self.assertEqual(CheckinRecord.objects.filter(member=self.author_member,checkin_date=timezone.localdate()).count(),1)
        self.assertEqual(CurrencyBalance.objects.get(member=self.author_member).balance,3)
        self.assertEqual(CurrencyTransaction.objects.filter(member=self.author_member,transaction_type='checkin').count(),1)

    def test_closed_current_week_does_not_display_a_future_week(self):
        self.enable();self.submit()
        self.week.is_closed=True;self.week.save()
        dashboard=self.client.get('/api/dashboard').json()
        self.assertEqual(dashboard['current_week']['id'],self.week.pk)
        self.assertTrue(dashboard['current_week']['closed'])
        members=self.client.get('/api/members').json()['members']
        self.assertEqual(next(m['submissions'] for m in members if m['id']==self.author_member.pk),1)
        self.assertEqual(self.client.get('/api/showcase').json()['week'],self.week.pk)

    def test_pending_state_without_timestamp_is_a_failure_not_an_exception(self):
        ContestSyncState.objects.create(user=self.author,status='pending')
        self.assertEqual(self.client.get('/api/contests/progress').json()['status'],'failed')

    def test_invalid_parameters_and_virtual_admin_mutations(self):
        self.enable()
        for path in ['/api/currency/transactions?limit=abc','/api/currency/transactions?limit=-1','/api/currency/checkin/calendar?month=0','/api/currency/checkin/calendar?year=bad']:
            self.assertEqual(self.client.get(path).status_code,400)
        self.client.force_login(User.objects.get(username='admin'))
        session=self.client.session;session['team_id']=self.team.pk;session.save()
        self.assertNotEqual(self.client.post('/api/currency/checkin').status_code,200)
        self.assertEqual(self.client.post('/api/currency/contest/submit','{}',content_type='application/json').status_code,403)

    def test_purchased_item_delete_retype_and_bad_shop_config_are_explicit(self):
        self.enable()
        item=ShopItem.objects.create(team=self.team,name='券',item_type='makeup_ticket',price=3)
        CurrencyBalance.objects.create(team=self.team,member=self.author_member,balance=20)
        purchase_item(self.author_member,item.pk)
        self.client.force_login(User.objects.get(username='admin'))
        session=self.client.session;session['team_id']=self.team.pk;session.save()
        endpoint='/api/currency/admin/shop'
        self.assertEqual(self.client.delete(endpoint,json.dumps({'item_id':item.pk}),content_type='application/json').status_code,409)
        self.assertEqual(self.client.put(endpoint,json.dumps({'item_id':item.pk,'item_type':'general'}),content_type='application/json').status_code,409)
        for data in [{'name':'Bad','price':'x'},{'name':'Bad','price':-1},{'name':'Bad','price':0},{'name':'Bad','stock':-2},{'name':'Bad','is_active':'false'},{'name':'Bad','icon':'not-an-image'}]:
            self.assertEqual(self.client.post(endpoint,json.dumps(data),content_type='application/json').status_code,400)
        response=self.client.put(endpoint,json.dumps({'item_id':item.pk,'is_active':False}),content_type='application/json')
        self.assertEqual(response.status_code,200)
        self.assertEqual(UserInventory.objects.get(member=self.author_member,item=item).quantity,1)

    def test_confirmed_team_deletion_with_purchase_history(self):
        self.enable()
        item=ShopItem.objects.create(team=self.team,name='Ticket',item_type='makeup_ticket',price=3)
        CurrencyBalance.objects.create(team=self.team,member=self.author_member,balance=20)
        purchase_item(self.author_member,item.pk)
        self.client.force_login(User.objects.get(username='admin'))
        response=self.client.delete(f'/api/management/teams/{self.team.pk}',json.dumps({'confirmation':self.team.code}),content_type='application/json')
        self.assertEqual(response.status_code,200,response.content)
        self.assertFalse(Team.objects.filter(pk=self.team.pk).exists())
        self.assertTrue(User.objects.filter(pk=self.author.pk).exists())

    def test_password_whitespace_and_user_cascade_delete(self):
        payload={'account':'audit-spaces','password':' pass-12345 ','confirm_password':' pass-12345 ','grade':0}
        self.client.logout()
        self.assertEqual(self.client.post('/api/auth/register',json.dumps(payload),content_type='application/json').status_code,201)
        self.client.post('/api/auth/logout')
        self.assertEqual(self.client.post('/api/auth/login',json.dumps({'account':'audit-spaces','password':' pass-12345 '}),content_type='application/json').status_code,200)
        self.assertEqual(self.client.get('/api/auth/logout').status_code,405)
        User.objects.get(username='audit-spaces').delete()
        self.author.delete()
        self.assertFalse(TeamMembership.objects.filter(pk=self.author_member.pk).exists())

    def test_purchase_request_retries_do_not_spend_twice(self):
        import uuid
        self.enable()
        item=ShopItem.objects.create(team=self.team,name='Ticket',item_type='makeup_ticket',price=3)
        CurrencyBalance.objects.create(team=self.team,member=self.author_member,balance=20)
        token=str(uuid.uuid4())
        first=purchase_item(self.author_member,item.pk,request_id=token)
        duplicate=purchase_item(self.author_member,item.pk,request_id=token)
        self.assertTrue(duplicate['duplicate'])
        self.assertEqual(CurrencyBalance.objects.get(member=self.author_member).balance,17)
        self.assertEqual(UserInventory.objects.get(member=self.author_member,item=item).quantity,1)
        self.assertEqual(Purchase.objects.filter(member=self.author_member).count(),1)
        with self.assertRaises(ValueError):purchase_item(self.author_member,item.pk,quantity=2,request_id=token)

    def test_sync_timeout_tracks_last_progress_instead_of_initial_request(self):
        state=ContestSyncState.objects.create(user=self.author,status='pending',remaining=20,total=30,
            requested_at=timezone.now()-timedelta(hours=2),progressed_at=timezone.now())
        self.assertEqual(self.client.get('/api/contests/progress').json()['status'],'pending')
        with patch('club.tasks.discover_contest_progress.delay') as dispatch:
            self.assertEqual(self.client.post('/api/contests/progress').status_code,202)
            dispatch.assert_not_called()
        state.progressed_at=timezone.now()-timedelta(minutes=31);state.save()
        self.assertEqual(self.client.get('/api/contests/progress').json()['status'],'failed')

    def test_duplicate_hydration_delivery_cannot_finish_other_contests(self):
        from .tasks import hydrate_contest_progress,enqueue_hydration
        state=ContestSyncState.objects.create(user=self.author,status='pending',remaining=2,total=2)
        c=ContestParticipation.objects.create(user=self.author,platform='codeforces',account='reader',
            contest_key='1',title='Fixture',url='https://codeforces.com/contest/1',starts_at=REWARD_START,ends_at=REWARD_START+timedelta(hours=2))
        with patch('club.contest_progress_sources.hydrate',return_value=([],'')) as hydrate:
            hydrate_contest_progress(c.pk,str(state.token))
            hydrate_contest_progress(c.pk,str(state.token))
        self.assertEqual(hydrate.call_count,1)
        state.refresh_from_db();self.assertEqual(state.remaining,1);self.assertEqual(state.status,'pending')
        with patch('club.tasks.hydrate_contest_progress.apply_async',side_effect=RuntimeError('broker down')):
            enqueue_hydration(c.pk,str(state.token))
        state.refresh_from_db();self.assertEqual(state.status,'failed')

    def test_partial_error_roster_preserves_previously_verified_ac(self):
        from .tasks import hydrate_contest_progress
        from .models import PlatformBinding,ContestProblemProgress
        state=ContestSyncState.objects.create(user=self.author,status='pending',remaining=1,total=1)
        PlatformBinding.objects.create(user=self.author,platform='atcoder',account='reader')
        c=ContestParticipation.objects.create(user=self.author,platform='atcoder',account='reader',contest_key='abc1',
            title='Fixture',url='https://atcoder.jp/contests/abc1',starts_at=REWARD_START,ends_at=REWARD_START+timedelta(hours=2),detail_complete=True)
        p=ContestProblemProgress.objects.create(participation=c,problem_key='a',label='A',title='One',accepted=True,official_ac=True,accepted_at=REWARD_START+timedelta(minutes=1),difficulty=400)
        with patch('club.contest_progress_sources.hydrate',return_value=([{'problem_key':'a','label':'A','title':'One','accepted':False,'official_ac':False,'accepted_at':None}], 'mirror unavailable')):
            hydrate_contest_progress(c.pk,str(state.token))
        p.refresh_from_db();c.refresh_from_db()
        self.assertTrue(p.official_ac);self.assertTrue(c.detail_complete)
        self.assertTrue(c.sync_error)

    def test_export_time_is_beijing_and_compatibility_models_are_identical(self):
        import io
        from openpyxl import load_workbook
        from . import currency_models,models
        self.assertIs(currency_models.CurrencyBalance,models.CurrencyBalance)
        item=self.submit()
        self.client.force_login(self.owner)
        response=self.client.get('/api/export')
        workbook=load_workbook(io.BytesIO(b''.join(response.streaming_content)))
        actual=workbook.active.cell(2,4).value
        self.assertEqual(actual,timezone.localtime(item.submitted_at).strftime('%Y-%m-%d %H:%M'))


class AuditDifficultyTests(SimpleTestCase):
    def test_each_platform_grade_thresholds_and_configured_amounts(self):
        cases={'leetcode':[(599,1,1),(600,2,1),(1000,2,2),(1200,3,2),(1500,3,3)],
               'codeforces':[(999,1,1),(1000,2,1),(1400,2,2),(1600,3,2),(2000,3,3)],
               'atcoder':[(399,1,1),(400,2,1),(800,2,2),(1000,3,2),(1600,3,3)],
               'nowcoder':[(1099,1,1),(1100,2,1),(1600,3,2),(2100,3,3)],
               'luogu':[(1,1,1),(2,2,1),(3,2,1),(4,3,2),(6,3,3)]}
        team=SimpleNamespace(currency_enabled=True,problem_rewards={},pk=1)
        member=SimpleNamespace(team=team,team_id=1)
        c=SimpleNamespace(source_data={},starts_at=REWARD_START,ends_at=REWARD_START+timedelta(hours=2),detail_complete=True)
        p=SimpleNamespace(participation=c,accepted=True,official_ac=True,accepted_at=REWARD_START+timedelta(minutes=20))
        with patch('django.utils.timezone.now',return_value=REWARD_START+timedelta(days=1)):
            for platform,rows in cases.items():
                c.platform=platform
                for difficulty,beginner,experienced in rows:
                    p.difficulty=difficulty
                    for grade,expected in [(0,beginner),(1,experienced),(4,experienced)]:
                        profile=SimpleNamespace(grade=grade,grade_date=date(2026,10,29))
                        self.assertEqual(reward_for(p,member,profile)[0],expected,(platform,difficulty,grade))
            c.platform='leetcode';p.difficulty=600
            team.problem_rewards={'beginner':[2,4,6],'experienced':[1,3,5]}
            self.assertEqual(reward_for(p,member,SimpleNamespace(grade=0,grade_date=date(2026,10,29)))[0],4)
            self.assertEqual(reward_for(p,member,SimpleNamespace(grade=1,grade_date=date(2026,10,29)))[0],1)
            p.difficulty=float('nan')
            self.assertEqual(reward_for(p,member,SimpleNamespace(grade=0,grade_date=date(2026,10,29)))[0],2)
            team.problem_rewards={'beginner':[]}
            self.assertEqual(reward_for(p,member,SimpleNamespace(grade=0,grade_date=date(2026,10,29)))[0],0)

    def test_unknown_time_fallback_for_all_platforms_but_known_outside_stays_blocked(self):
        c=SimpleNamespace(source_data={},starts_at=REWARD_START,ends_at=REWARD_START+timedelta(hours=2),detail_complete=False)
        member=SimpleNamespace(team=SimpleNamespace(currency_enabled=True,problem_rewards={}),team_id=1)
        profile=SimpleNamespace(grade=0,grade_date=date(2026,10,29))
        p=SimpleNamespace(participation=c,accepted=True,official_ac=False,accepted_at=None,difficulty=None)
        with patch('django.utils.timezone.now',return_value=REWARD_START+timedelta(days=1)):
            for platform in ['codeforces','atcoder','nowcoder','luogu','leetcode']:
                c.platform=platform
                self.assertEqual(reward_for(p,member,profile)[0],1)
                p.accepted_at=c.ends_at;c.detail_complete=True
                self.assertEqual(reward_for(p,member,profile)[0],0)
                p.accepted_at=None;p.accepted=False
                self.assertEqual(reward_for(p,member,profile)[0],0)
                p.accepted=True;c.detail_complete=False

    def test_luogu_time_endpoint_failure_uses_reliable_passed_list(self):
        from .contest_progress_sources import hydrate,SyncError
        remote=Mock()
        remote.json.side_effect=[{'data':{'passed':['P1']}},{'data':{'contestProblems':[{'no':'A','problem':{'pid':'P1','name':'One','difficulty':2}}]}}]
        c=SimpleNamespace(platform='luogu',contest_key='1',account='9',url='https://www.luogu.com.cn/contest/1',starts_at=REWARD_START,ends_at=REWARD_START+timedelta(hours=2))
        with patch('club.contest_progress_sources.configure_luogu',return_value='cookie'),patch('club.platform_sync.luogu',return_value={'profile_url':'https://www.luogu.com.cn/user/9'}),patch('club.contest_progress_sources.luogu_ac_times',side_effect=SyncError('temporary denial')):
            rows,error=hydrate(remote,c)
        self.assertTrue(rows[0]['accepted']);self.assertIsNone(rows[0]['accepted_at']);self.assertFalse(error)

    def test_known_outside_nowcoder_time_does_not_become_unknown(self):
        from .contest_progress_sources import hydrate
        remote=Mock();start=REWARD_START.timestamp()
        c=SimpleNamespace(platform='nowcoder',contest_key='88',account='9',url='https://ac.nowcoder.com/acm/contest/88',starts_at=REWARD_START,ends_at=REWARD_START+timedelta(hours=2))
        remote.json.return_value={'data':{'data':[{'index':'A','problemId':1,'title':'One'}]}}
        with patch('club.contest_progress_sources.nowcoder_accepted',return_value={'problem/1':[start-10000]}),patch('club.contest_progress_sources.nowcoder_ratings',return_value={'1':1600}):rows,error=hydrate(remote,c)
        self.assertTrue(rows[0]['accepted']);self.assertIsNotNone(rows[0]['accepted_at']);self.assertFalse(rows[0]['official_ac'])

    def test_nowcoder_short_server_pages_and_negative_unknown_difficulty(self):
        from .contest_progress_sources import nowcoder_ratings
        remote=Mock()
        remote.json.side_effect=[{'code':0,'data':{'totalCount':2,'papers':[{'questions':[{'problemId':1,'difficulty':1600}]}]}},
            {'code':0,'data':{'totalCount':2,'papers':[{'questions':[{'problemId':2,'difficulty':2100},{'problemId':3,'difficulty':-1}]}]}}]
        with patch('club.contest_progress_sources.cache.get',return_value=None):
            result=nowcoder_ratings(remote)
        self.assertEqual(result,{'1':1600,'2':2100})
        self.assertEqual(remote.json.call_count,2)

    def test_empty_atcoder_mirror_is_incomplete_instead_of_verified_zero_ac(self):
        from .contest_progress_sources import hydrate
        remote=Mock();remote.text.return_value='<table><tbody><tr><td><a href="/contests/abc1/tasks/abc1_a">A</a></td><td>One</td></tr></tbody></table>'
        remote.json.return_value=[]
        c=SimpleNamespace(platform='atcoder',contest_key='abc1',account='reader',url='https://atcoder.jp/contests/abc1',starts_at=REWARD_START,ends_at=REWARD_START+timedelta(hours=2))
        with patch('club.contest_progress_sources.cached',return_value={'abc1_a':{'difficulty':400}}):rows,error=hydrate(remote,c)
        self.assertEqual(len(rows),1);self.assertEqual(rows[0]['difficulty'],400)
        self.assertTrue(error)
        self.assertFalse(rows[0]['official_ac'])

    def test_luogu_current_schema_labels_difficulty_and_official_ac(self):
        from .contest_progress_sources import hydrate
        remote=Mock();start=REWARD_START.timestamp()
        contest=SimpleNamespace(platform='luogu',contest_key='88',account='9',url='https://www.luogu.com.cn/contest/88',starts_at=REWARD_START,ends_at=REWARD_START+timedelta(hours=2))
        remote.json.side_effect=[{'data':{'passed':['P1','P2']}},{'data':{'contestProblems':[
            {'no':'T1','problem':{'pid':'P1','name':'One','difficulty':4}}, {'no':'T2','problem':{'pid':'P2','name':'Two','difficulty':0}}]}}]
        with patch('club.contest_progress_sources.configure_luogu',return_value='cookie'),patch('club.platform_sync.luogu',return_value={'profile_url':'https://www.luogu.com.cn/user/9'}),patch('club.contest_progress_sources.luogu_ac_times',return_value={'P1':[start+10],'P2':[start+10000]}):
            rows,error=hydrate(remote,contest)
        self.assertEqual([r['label'] for r in rows],['T1','T2'])
        self.assertEqual(rows[0]['difficulty'],4)
        self.assertTrue(rows[0]['official_ac']);self.assertFalse(rows[1]['official_ac'])
        self.assertIsNone(rows[1]['difficulty'])

    def test_luogu_practice_ac_during_contest_is_not_official(self):
        from .contest_progress_sources import luogu_ac_times
        remote=Mock();remote.json.side_effect=[{'data':{'records':{'result':[
            {'id':1,'status':12,'submitTime':1000,'problem':{'pid':'P1'},'user':{'uid':9},'contest':None},
            {'id':2,'status':12,'submitTime':1001,'problem':{'pid':'P2'},'user':{'uid':9},'contest':{'id':88}}]}}},
            {'data':{'records':{'result':[]}}}]
        with patch('club.contest_progress_sources.time.sleep'):
            result=luogu_ac_times(remote,'9','_uid=1; __client_id=secret',0,'88')
        self.assertEqual(result,{'P2':[1001]})

    def test_nowcoder_follows_small_html_pages_and_retains_multiple_times(self):
        from .contest_progress_sources import nowcoder_accepted
        remote=Mock()
        def html(day,next_page=False):
            return '<script>window.curUser.id="9";</script><table><tr><td>1</td><td><a href="/acm/problem/1">One</a></td><td>答案正确</td><td>'+day+'</td></tr></table>'+('<div class="pagination"><a href="?page=2">下一页</a></div>' if next_page else '')
        remote.text.side_effect=[html('2026-11-01 12:00:00',True),html('2026-10-31 12:00:00')]
        with patch('club.contest_progress_sources.cache.get',return_value=None):result=nowcoder_accepted(remote,'9')
        self.assertEqual(len(result['problem/1']),2)
        self.assertEqual(remote.text.call_count,2)

    def test_nowcoder_multiple_ac_keeps_in_contest_evidence(self):
        from .contest_progress_sources import hydrate
        start=REWARD_START.timestamp();remote=Mock()
        c=SimpleNamespace(platform='nowcoder',contest_key='88',account='9',url='https://ac.nowcoder.com/acm/contest/88',starts_at=REWARD_START,ends_at=REWARD_START+timedelta(hours=2))
        remote.json.return_value={'data':{'data':[{'index':'A','problemId':1,'title':'One'}]}}
        with patch('club.contest_progress_sources.nowcoder_accepted',return_value={'problem/1':[start-10000,start+50]}),patch('club.contest_progress_sources.nowcoder_ratings',return_value={'1':1600}):
            rows,error=hydrate(remote,c)
        self.assertTrue(rows[0]['official_ac'])
        self.assertEqual(rows[0]['accepted_at'].timestamp(),start+50)


class AuditConcurrentWeekTests(TransactionTestCase):
    def test_simultaneous_weekly_settlement_only_credits_once(self):
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier
        root=User.objects.filter(username='admin').first() or User.objects.create_superuser('admin',password='audit-test-only')
        user=User.objects.create_user('audit-concurrent')
        team=Team.objects.create(name='Audit concurrent',code='AUDIT-CONCURRENT',created_by=root,currency_enabled=True)
        member=TeamMembership.objects.create(team=team,user=user,display_name='Concurrent')
        monday=date(2026,10,19)
        StreakReward.objects.create(team=team,days=3,reward=2)
        for offset in range(3):CheckinRecord.objects.create(team=team,member=member,checkin_date=monday+timedelta(days=offset))
        barrier=Barrier(2)
        def run(_):
            close_old_connections()
            try:
                local=TeamMembership.objects.get(pk=member.pk)
                barrier.wait(timeout=10)
                return settle_weekly_streak(local,monday)
            finally:close_old_connections()
        with ThreadPoolExecutor(2) as pool:results=list(pool.map(run,[1,2]))
        self.assertEqual(sorted(results),[0,2])
        self.assertEqual(CurrencyBalance.objects.get(member=member).balance,2)
        self.assertEqual(WeeklyStreakGrant.objects.filter(member=member).count(),1)


    def test_purchase_and_makeup_concurrency_preserves_inventory(self):
        from concurrent.futures import ThreadPoolExecutor
        from threading import Barrier
        from .models import Semester,MakeupPermission
        from .currency_services import makeup_checkin
        from .services import build_weeks
        root=User.objects.filter(username='admin').first() or User.objects.create_superuser('admin',password='audit-test-only')
        user=User.objects.create_user('audit-inventory')
        team=Team.objects.create(name='Inventory race',code='AUDIT-INVENTORY',created_by=root,currency_enabled=True)
        member=TeamMembership.objects.create(team=team,user=user,display_name='Inventory')
        today=timezone.localdate()
        semester=Semester.objects.create(team=team,name='Current',starts_on=today-timedelta(days=10),ends_on=today+timedelta(days=10))
        build_weeks(semester)
        item=ShopItem.objects.create(team=team,name='Ticket',item_type='makeup_ticket',price=3)
        UserInventory.objects.create(team=team,member=member,item=item,quantity=1)
        CurrencyBalance.objects.create(team=team,member=member,balance=20)
        barrier=Barrier(2)
        def run(operation):
            close_old_connections()
            try:
                local=TeamMembership.objects.get(pk=member.pk)
                barrier.wait(timeout=10)
                return purchase_item(local,item.pk) if operation=='buy' else makeup_checkin(local,today-timedelta(days=1))
            finally:close_old_connections()
        with ThreadPoolExecutor(2) as pool:results=list(pool.map(run,['buy','makeup']))
        self.assertTrue(all(r['success'] for r in results))
        self.assertEqual(UserInventory.objects.get(member=member,item=item).quantity,1)
        self.assertEqual(CurrencyBalance.objects.get(member=member).balance,17)
        self.assertEqual(MakeupPermission.objects.filter(member=member).count(),1)
