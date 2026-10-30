import json
import io
import warnings
from urllib.parse import quote

from PIL import Image, UnidentifiedImageError
from django.conf import settings
from django.core.exceptions import ValidationError
from django.core.validators import URLValidator
from django.http import HttpResponse
from datetime import datetime, timedelta
from pathlib import Path

from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.models import User
from django.db import IntegrityError, transaction
from django.db.models import Count, Q, Sum
from django.http import FileResponse, Http404, JsonResponse
from django.middleware.csrf import get_token
from django.utils import timezone
from django.views.decorators.csrf import ensure_csrf_cookie
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from django.contrib.auth.hashers import check_password
from functools import wraps

from .models import FileDeletion, ReviewLog, Semester, SemesterMemberTotal, Submission, SubmissionPart, Team, TeamMembership, Week, create_internal_username
from .services import build_weeks, progress_rows, qualified, refresh_totals, week_bounds, score_part

BASE_DIR = Path(__file__).resolve().parent.parent


def api_error(message, status=400):
    return JsonResponse({"error": message}, status=status)


def api_login_required(view):
    """Return JSON 401 responses for API calls instead of HTML redirects."""
    @wraps(view)
    def wrapped(request, *args, **kwargs):
        if not request.user.is_authenticated:
            return api_error("请先登录", 401)
        return view(request, *args, **kwargs)
    return wrapped


@ensure_csrf_cookie
def api_csrf(request):
    if request.method != "GET":
        return api_error("仅支持 GET", 405)
    response = JsonResponse({"csrfToken": get_token(request)})
    response["Cache-Control"] = "no-store"
    return response


def body(request):
    try:
        data = json.loads(request.body or "{}")
        return data if isinstance(data, dict) else {}
    except (json.JSONDecodeError, UnicodeDecodeError):
        return {}


def membership_for(request, team_id=None):
    if not request.user.is_authenticated:
        return None
    team_id = team_id or request.session.get("team_id")
    if request.user.is_superuser:
        team = Team.objects.filter(pk=team_id).first() if team_id else Team.objects.order_by('id').first()
        if not team:
            return None
        # Administrators can inspect/manage teams without joining or affecting rankings.
        return TeamMembership.objects.filter(user=request.user, team=team, is_active=True).first() or TeamMembership(
            id=0, user=request.user, team=team, display_name='管理员', role=TeamMembership.ADMIN)
    queryset = TeamMembership.objects.select_related("team").filter(user=request.user, is_active=True)
    return queryset.filter(team_id=team_id).first() if team_id else queryset.first()


def teams_for(user):
    if user.is_superuser:
        return [{'team_id': t.pk, 'team__name': t.name, 'team__code': t.code} for t in Team.objects.order_by('id')]
    return list(user.team_memberships.filter(is_active=True).values('team_id', 'team__name', 'team__code'))


def active_semester(team):
    from .semester_calendar import ensure_current_semester
    ensure_current_semester(team)
    today = timezone.localdate()
    semester = team.semesters.filter(is_active=True, archived_at__isnull=True, starts_on__lte=today, ends_on__gte=today).first()
    if semester:
        build_weeks(semester)
    return semester


def active_week(semester):
    now = timezone.now()
    # This helper is used by read-only dashboards/lists; closed weeks still display.
    return semester.weeks.filter(starts_at__lte=now, ends_at__gt=now).first()


def member_week_score(member, week):
    return progress_rows(week.semester, week, [member.id])[0].get(member.id, 0)


def member_week_submissions(member, week):
    return progress_rows(week.semester, week, [member.id])[1].get(member.id, 0)


def week_progress_qualified(score, submissions, week):
    return qualified(score, submissions, week)


def member_week_qualified(member, week):
    scores, counts = progress_rows(week.semester, week, [member.id])
    return qualified(scores.get(member.id, 0), counts.get(member.id, 0), week)


def member_week_progress(members, week):
    return progress_rows(week.semester, week, [member.id for member in members]) if week and members else ({}, {})


def leaderboard_rows_for(team, semester, week=None):
    scores, counts = progress_rows(semester, week)
    rows = [{'memberId': member.id, 'name': member.display_name, 'score': scores.get(member.id, 0),
             'submissions': counts.get(member.id, 0)} for member in team.memberships.filter(is_active=True,user__is_active=True,user__is_superuser=False)]
    rows.sort(key=lambda row: (-row['score'], -row['submissions'], row['name']))
    for index, row in enumerate(rows, 1):
        row['rank'] = index
    return rows


def member_streak(member, semester, through_week, *, progress=None):
    weeks = list(semester.weeks.filter(number__lte=through_week.number).order_by("-number"))
    scores, counts = progress if progress is not None else progress_rows(semester, member_ids=[member.id], by_week=True)
    streak = 0
    for week in weeks:
        # The open week is still in progress and cannot be counted as a
        # completed streak until the scheduled settlement task closes it.
        if week.id == through_week.id and not week.is_closed:
            continue
        key = (week.id, member.id)
        if not qualified(scores.get(key, 0), counts.get(key, 0), week):
            break
        streak += 1
    return streak


def serialize_part(part, include_content=True):
    payload = {"id": part.id, "kind": part.kind, "status": part.status, "points": part.points_awarded, "note": part.review_note}
    if include_content:
        payload.update({"text": part.text_content, "link": part.link, "upload": f"/api/attachments/{part.id}" if part.upload else None})
    return payload


def serialize_submission(submission, include_content=True, parts=None):
    selected_parts = parts if parts is not None else submission.parts.all()
    return {"id": submission.id, "member_id": submission.member_id, "member_name": submission.member.display_name, "week": submission.week.number, "week_id": submission.week_id, "submitted_at": submission.submitted_at.isoformat(), "record_date": submission.record_date.isoformat(), "is_valid": submission.is_valid, "version": submission.version, "rejection_note": submission.rejection_note, "problem_url": submission.problem.url if submission.problem_id else "", "problem_title": (submission.problem.title or submission.problem.url) if submission.problem_id else "", "semester_name": submission.semester.name, "can_delete": not bool(submission.semester.archived_at), "can_edit": not submission.semester.archived_at and submission.semester.is_active and not submission.week.is_closed and submission.week.starts_at <= timezone.now() < submission.week.ends_at, "can_resubmit": not submission.is_valid and not submission.semester.archived_at and not submission.week.is_closed and submission.week.starts_at <= timezone.now() < submission.week.ends_at, "parts": [serialize_part(part, include_content) for part in selected_parts]}


@ensure_csrf_cookie
def frontend(request):
    path = BASE_DIR / "index.html"
    if not path.exists():
        raise Http404
    get_token(request)
    raw = repair_text(path.read_text(encoding="utf-8-sig"))
    from django.http import HttpResponse
    response = HttpResponse(raw, content_type="text/html; charset=utf-8")
    response["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response["Pragma"] = "no-cache"
    return response


@ensure_csrf_cookie
def auth_page(request):
    path = BASE_DIR / "auth.html"
    if not path.exists():
        raise Http404
    get_token(request)
    from django.http import HttpResponse
    response = HttpResponse(repair_text(path.read_text(encoding="utf-8-sig")), content_type="text/html; charset=utf-8")
    response["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response["Pragma"] = "no-cache"
    return response


def repair_text(value):
    import re
    def fix(match):
        chunk = match.group(0)
        try:
            return chunk.encode("latin1").decode("utf-8")
        except UnicodeError:
            return chunk
    return re.sub(r"[^\x00-\x7f]+", fix, value)


@api_login_required
@transaction.atomic
def join_team(request):
    if request.method != 'POST':
        return api_error('仅支持POST', 405)
    if request.user.is_superuser:return api_error('系统管理员可直接切换管理团队，无需加入',403)
    User.objects.select_for_update(no_key=True).get(pk=request.user.pk)
    data = body(request)
    code,name=[str(data.get(key,'')).strip() for key in ('team_code','name')]
    if not code or len(code) > 32 or len(name) > 40:
        return api_error('请填写有效团队码，团队昵称不能超过40字')
    team = Team.objects.filter(code__iexact=code).first()
    if not team:
        return api_error('团队码不存在，请检查后重试', 404)
    if not name:
        prior = TeamMembership.objects.filter(user=request.user, is_active=True).order_by('-joined_at').first()
        base = (prior.display_name if prior else request.user.first_name or request.user.username)[:40]
        name = base
        suffix = 2
        while TeamMembership.objects.filter(team=team, display_name=name).exists():
            tail = f'-{suffix}'
            name = base[:40-len(tail)] + tail
            suffix += 1
    if TeamMembership.objects.filter(team=team).filter(Q(user=request.user) | Q(display_name=name)).exists():
        return api_error('已经加入该团队或昵称已被使用')
    try:
        with transaction.atomic():
            membership = TeamMembership.objects.create(user=request.user, team=team, display_name=name, role=TeamMembership.MEMBER)
    except IntegrityError:
        return api_error('昵称已被使用或团队信息已变化，请刷新后重试',409)
    request.session['team_id'] = team.pk
    return JsonResponse({'ok': True, 'team': team.name, 'name': membership.display_name})


def register_account(request):
    if request.method != "POST":
        return api_error("仅支持 POST", 405)
    data = body(request)
    account=data.get('account','');password=data.get('password','');confirm_password=data.get('confirm_password','')
    if not all(isinstance(v,str) for v in (account,password,confirm_password)):
        return api_error('账号或密码格式无效')
    account=account.strip()
    if not 3 <= len(account) <= 150 or not 8 <= len(password) <= 128:
        return api_error('账号需为3–150位，密码需为8–128位')
    if password != confirm_password:
        return api_error("两次密码不一致")
    if User.objects.filter(username=account).exists():
        return api_error("该账号已存在")
    from .grades import validate_grade
    from .models import UserProfile
    try:
        grade = validate_grade(data.get('grade'))
    except ValueError as exc:
        return api_error(str(exc))
    try:
        with transaction.atomic():
            user = User.objects.create_user(username=account, password=password)
            UserProfile.objects.create(user=user, grade=grade, grade_date=timezone.localdate())
    except IntegrityError:
        return api_error('该账号已存在',409)
    login(request, user)
    request.session.pop("team_id", None)
    return JsonResponse({"ok": True, "account": user.username, "teams": []}, status=201)


def login_view(request):
    if request.method != "POST":
        return api_error("仅支持 POST", 405)
    data = body(request)
    code = str(data.get("team_code", "")).strip()
    name = str(data.get("name", "")).strip()
    account = str(data.get("account", name)).strip()
    password = str(data.get("password", ""))
    account_user = User.objects.filter(username=account, is_active=True).first()
    if account_user and check_password(password, account_user.password):
        memberships = TeamMembership.objects.select_related("team").filter(user=account_user, is_active=True)
        if code: memberships = memberships.filter(team__code__iexact=code)
        membership = memberships.filter(role=TeamMembership.ADMIN).first() or memberships.first()
        login(request, account_user)
        if membership:
            request.session["team_id"] = membership.team_id
        else:
            request.session.pop("team_id", None)
        return JsonResponse({"ok": True, "team": membership.team.name if membership else None, "name": membership.display_name if membership else account_user.username, "role": "site_admin" if account_user.is_superuser else membership.role if membership else "member", "site_admin": account_user.is_superuser, "teams": teams_for(account_user)})
    return api_error("账号或密码错误", 403)


def logout_view(request):
    if request.method != "POST":return api_error("仅支持 POST",405)
    logout(request)
    return JsonResponse({"ok": True})


@api_login_required
@transaction.atomic
def create_team(request):
    if request.method != "POST":
        return api_error("仅支持 POST", 405)
    if not request.user.is_superuser:
        return api_error("仅管理员可以创建团队", 403)
    data = body(request)
    name = str(data.get("team_name", "")).strip()
    code = str(data.get("team_code", "")).strip()
    if not name or not code:
        return api_error("请填写团队名称和团队码")
    if Team.objects.filter(code__iexact=code).exists():
        return api_error("团队编号已存在")
    membership = TeamMembership.objects.filter(user=request.user).first()
    display_name = membership.display_name if membership else str(data.get("display_name", request.user.username)).strip()
    if not display_name:
        return api_error("请填写创建者姓名")
    if len(name) > 80 or len(code) > 32 or len(display_name) > 40:
        return api_error('团队名称、编号或成员姓名过长')
    try:
        numeric = {}
        for key, default, upper in [('min_score', 10, 100000), ('min_submissions', 4, 100000), ('screenshot_points', 1, 10000), ('logic_points', 1, 10000), ('blog_points', 1, 10000)]:
            value = data.get(key, default)
            if isinstance(value, bool) or str(value) != str(int(value)) or not 0 <= int(value) <= upper:
                raise ValueError()
            numeric[key] = int(value)
    except (ValueError, TypeError):
        return api_error('积分设置无效')
    if type(data.get('require_both', True)) is not bool:
        return api_error('达标规则无效')
    team = Team.objects.create(name=name, code=code, created_by=request.user)
    from .semester_calendar import semester_bounds, semester_name
    start, end = semester_bounds(timezone.localdate())
    semester = Semester.objects.create(team=team, name=semester_name(start), starts_on=start, ends_on=end, calendar_managed=True, min_score=numeric["min_score"], min_submissions=numeric["min_submissions"], require_both=data.get("require_both", True), screenshot_points=numeric["screenshot_points"], logic_points=numeric["logic_points"], blog_points=numeric["blog_points"])
    build_weeks(semester)
    request.session["team_id"] = team.id
    return JsonResponse({"ok": True, "team": {"id": team.id, "name": team.name, "code": team.code}, "name": display_name})


@api_login_required
def select_team(request):
    if request.method != "POST": return api_error("仅支持 POST", 405)
    team_id = body(request).get("team_id")
    if not (request.user.is_superuser and Team.objects.filter(pk=team_id).exists()) and not TeamMembership.objects.filter(user=request.user, team_id=team_id, is_active=True).exists(): return api_error("无权进入该团队", 403)
    request.session["team_id"] = team_id
    return JsonResponse({"ok": True})


@api_login_required
def me(request):
    from .profile_views import avatar_url
    membership = membership_for(request)
    teams = teams_for(request.user)
    role = 'site_admin' if request.user.is_superuser else membership.role if membership else 'member'
    return JsonResponse({'name': membership.display_name if membership else request.user.username,
        'account': request.user.username, 'avatar': avatar_url(request.user), 'site_admin': request.user.is_superuser, 'role': role,
        'is_team_manager': request.user.team_memberships.filter(role='admin', is_active=True).exists(),
        'can_manage_current': request.user.is_superuser or bool(membership and membership.is_admin),
        'managed_teams': list(request.user.team_memberships.filter(role='admin', is_active=True).values('team_id', 'team__name', 'team__code')),
        'can_submit': bool(membership and membership.pk and not request.user.is_superuser),
        'group': '管理员' if request.user.is_superuser else '团队管理用户组' if request.user.team_memberships.filter(role='admin', is_active=True).exists() else '普通用户组',
        'team': {'id': membership.team_id, 'name': membership.team.name, 'code': membership.team.code} if membership else None,
        'teams': teams, 'csrfToken': get_token(request)})


@api_login_required
def dashboard(request):
    membership = membership_for(request)
    if not membership:
        return api_error("未加入当前团队", 403)
    semester = active_semester(membership.team)
    if not semester:
        return JsonResponse({"team": {"id": membership.team_id, "name": membership.team.name}, "semester": None, "current_week": None, "stats": {"rank": None, "memberCount": 0, "streakWeeks": 0, "pendingParts": 0}})
    week = active_week(semester)
    if not week:
        week = semester.weeks.order_by("-number").first()
    parts = SubmissionPart.objects.filter(submission__member=membership, submission__week=week, submission__is_valid=True, status=SubmissionPart.APPROVED)
    score = parts.aggregate(total=Sum("points_awarded"))["total"] or 0
    count = member_week_submissions(membership, week)
    pending_parts = SubmissionPart.objects.filter(submission__member=membership, submission__week=week, status=SubmissionPart.PENDING).count()
    rows = leaderboard_rows_for(membership.team, semester, week)
    current_rank = next((row["rank"] for row in rows if row["memberId"] == membership.id), None)
    previous_week = semester.weeks.filter(number=week.number - 1).first()
    previous_rank = None
    if previous_week:
        previous_rows = leaderboard_rows_for(membership.team, semester, previous_week)
        previous_rank = next((row["rank"] for row in previous_rows if row["memberId"] == membership.id), None)
    now = timezone.now()
    semester_deadline = timezone.make_aware(datetime.combine(semester.ends_on + timedelta(days=1), datetime.min.time()))
    week_deadline = min(week.ends_at, semester_deadline)
    remaining_seconds = max(0, int((week_deadline - now).total_seconds())) if not week.is_closed else 0
    history_scores, history_counts = progress_rows(semester, member_ids=[membership.id], by_week=True)
    today = timezone.localdate(now)
    dates = set(Submission.objects.filter(member=membership, is_valid=True).values_list('record_date', flat=True))
    from .models import CheckinRecord, MakeupPermission
    permissions = {p.record_date: p for p in MakeupPermission.objects.filter(member=membership, semester=semester)}
    checkins = dict(CheckinRecord.objects.filter(member=membership).values_list('checkin_date', 'is_makeup'))
    submission_dates = dates.copy()
    dates.update(checkins)

    # Show seven past days, today, and the next three days.
    start = today - timedelta(days=7)
    weeks = list(semester.weeks.all())
    pulse_days = []
    for offset in range(11):
        day = start + timedelta(days=offset)
        day_week = next((w for w in weeks if timezone.localdate(w.starts_at) <= day < timezone.localdate(w.ends_at)), None)
        permission = permissions.get(day)
        completed = bool(checkins.get(day) or permission and permission.completed_at)
        can_submit = bool(membership.pk and day_week and semester.is_active and not semester.archived_at and semester.starts_on <= day <= min(today, semester.ends_on))
        if day == today:
            can_submit = can_submit and not day_week.is_closed and day_week.starts_at <= now < day_week.ends_at
        elif day < today:
            can_submit = can_submit and day not in dates and not completed

        pulse_days.append({
            'date': day.isoformat(),
            'submitted': day in dates,
            'hasSubmission': day in submission_dates,
            'checkedIn': day in checkins,
            'isMakeup': completed,
            'makeupOpen': bool(permission and not permission.completed_at),
            'weekId': day_week.pk if day_week else None,
            'isToday': day == today,
            'canSubmit': can_submit
        })

    cursor = today if today in dates else today - timedelta(days=1)
    streak = 0
    while cursor in dates:
        streak += 1
        cursor -= timedelta(days=1)
    return JsonResponse({
        "team": {"id": membership.team_id, "name": membership.team.name},
        "semester": {"id": semester.id, "name": semester.name, "minScore": week.effective_min_score, "minSubmissions": week.effective_min_submissions, "requireBoth": week.effective_require_both, "totalWeeks": semester.weeks.count()},
        "current_week": {"id": week.id, "number": week.number, "startsAt": week.starts_at.isoformat(), "endsAt": week_deadline.isoformat(), "closed": week.is_closed, "remainingSeconds": remaining_seconds},
        "score": score,
        "submissions": count,

        "stats": {"rank": current_rank, "previousRank": previous_rank, "memberCount": len(rows), "streakWeeks": member_streak(membership, semester, week, progress=(history_scores, history_counts)), "pendingParts": pending_parts},
        "pulse": pulse_days, "today": today.isoformat(), "streakDays": streak,
    })


@api_login_required
def members(request):
    membership = membership_for(request)
    if not membership:
        return api_error("未加入当前团队", 403)
    semester = active_semester(membership.team)
    week = active_week(semester) if semester else None
    scores, counts = progress_rows(semester, week) if week else ({}, {})
    from .models import ContestProblemProgress, PlatformBinding
    ac_counts = {}
    submission_ac_counts = {}
    if week:
        submission_ac_counts = dict(Submission.objects.filter(
            team=membership.team, week=week, is_valid=True, problem__isnull=False,
            parts__kind=SubmissionPart.PROOF, parts__status=SubmissionPart.APPROVED,
        ).order_by().values('member_id').annotate(total=Count('problem_id', distinct=True))
            .values_list('member_id', 'total'))
        user_ids = membership.team.memberships.filter(is_active=True, user__is_active=True,
                                                     user__is_superuser=False).values_list('user_id', flat=True)
        bindings = set(PlatformBinding.objects.filter(user_id__in=user_ids)
                       .values_list('user_id', 'platform', 'account'))
        accepted = ContestProblemProgress.objects.filter(
            participation__user_id__in=user_ids, accepted=True,
            participation__starts_at__gte=week.starts_at,
            participation__starts_at__lt=week.ends_at,
        ).values_list('participation__user_id', 'participation__platform', 'participation__account',
                      'participation__contest_key', 'problem_key').distinct()
        problems_by_user = {}
        for user_id, platform, account, contest_key, problem_key in accepted:
            if (user_id, platform, account) in bindings:
                problems_by_user.setdefault(user_id, set()).add((platform, contest_key, problem_key))
        ac_counts = {user_id: len(problems) for user_id, problems in problems_by_user.items()}
    rows = [{'id': member.id, 'name': member.display_name, 'role': member.role, 'team_id': membership.team_id,
             'account': member.user.username if request.user.is_superuser else '',
             'is_system_admin': member.user.is_superuser,
             'can_kick': not member.user.is_superuser and member.user_id != request.user.pk and
                         (request.user.is_superuser or (membership.is_admin and member.role == TeamMembership.MEMBER)),
             'can_manage_role': request.user.is_superuser and not member.user.is_superuser,
             'score': scores.get(member.id, 0), 'submissions': counts.get(member.id, 0),
             'ac_problems': ac_counts.get(member.user_id, 0) + submission_ac_counts.get(member.id, 0)}
            for member in membership.team.memberships.filter(is_active=True,user__is_active=True,user__is_superuser=False).select_related('user')]
    if request.user.is_superuser and membership.team.currency_enabled:
        from .models import CurrencyBalance
        balances = dict(CurrencyBalance.objects.filter(team=membership.team).values_list('member_id', 'balance'))
        for row in rows:
            row['currency_balance'] = balances.get(row['id'], 0)
    return JsonResponse({'members': rows, 'team_name': membership.team.name})


@api_login_required
def member_history(request, member_id):
    viewer = membership_for(request)
    if not viewer:
        return api_error("未加入当前团队", 403)
    target = viewer.team.memberships.filter(id=member_id).first()
    if not target:
        return api_error("成员不存在", 404)
    semester_id = request.GET.get('semester')
    if semester_id and not semester_id.isdigit():
        return api_error('学期编号无效')
    semester = (viewer.team.semesters.filter(pk=semester_id).first() if semester_id else
                active_semester(viewer.team) or viewer.team.semesters.order_by('-starts_on').first())
    if semester_id and not semester:
        return api_error('学期不存在', 404)
    if semester and semester.archived_at:
        total = semester.member_totals.filter(member=target).first()
        return JsonResponse({'member': {'id': target.id, 'name': total.display_name if total else target.display_name},
            'archived': True, 'semester': {'id': semester.id, 'name': semester.name},
            'summary': {'score': total.total_score if total else 0, 'submissions': total.total_submissions if total else 0,
                        'qualifiedWeeks': total.qualified_weeks if total else 0, 'qualified': None}, 'submissions': []})
    week_id = request.GET.get("week")
    if week_id and not week_id.isdigit():
        return api_error("周次编号无效")
    queryset = target.submissions.select_related("week").prefetch_related("parts").filter(team=viewer.team, semester=semester)
    if not viewer.is_admin and target.id != viewer.id:
        queryset = queryset.filter(is_valid=True)
    if week_id:
        queryset = queryset.filter(week__semester=semester, week__number=week_id)
    visible_submissions = list(queryset)
    summary_week = semester.weeks.filter(number=week_id).first() if week_id and semester else None
    scores, counts = progress_rows(semester, summary_week, [target.id], [item.id for item in visible_submissions])
    summary_score = scores.get(target.id, 0)
    summary_submissions = counts.get(target.id, 0)
    summary_qualified = qualified(summary_score, summary_submissions, summary_week) if summary_week else None
    return JsonResponse({"member": {"id": target.id, "name": target.display_name}, "summary": {"score": summary_score, "submissions": summary_submissions, "qualified": summary_qualified}, "submissions": [serialize_submission(item, True, parts=None if viewer.is_admin or target.id == viewer.id else [part for part in item.parts.all() if part.status == SubmissionPart.APPROVED]) for item in visible_submissions]})


@api_login_required
def submissions(request, submission_id=None, edit=False):
    membership = membership_for(request)
    if not membership:
        return api_error("未加入当前团队", 403)
    if request.method == "GET" and submission_id is None:
        queryset = membership.team.submissions.select_related("member", "week", "semester", "problem").prefetch_related("parts")
        if request.GET.get("mine") == "1":
            queryset = queryset.filter(member=membership)
        elif not membership.is_admin:
            queryset = queryset.filter(Q(member=membership) | Q(is_valid=True))
        week_id = request.GET.get("week")
        if week_id:
            if not week_id.isdigit():
                return api_error('周次编号无效')
            queryset = queryset.filter(week_id=week_id)
        return JsonResponse({'submissions': [serialize_submission(item) for item in queryset]})
    if request.method != "POST":
        return api_error("仅支持 GET 或 POST", 405)
    try:
        week_id = int(request.POST.get('week_id', ''))
    except (ValueError, TypeError):
        return api_error('请刷新页面，确认提交周次后重试')
    if not membership.pk or request.user.is_superuser:
        return api_error("请以团队成员身份提交训练记录", 403)
    from .problem_parser import canonical_problem, get_problem
    try:
        canonical_problem(request.POST.get('problem_url', ''))
    except ValidationError as exc:
        return api_error('; '.join(exc.messages))
    proof = request.FILES.get('proof')
    if not proof and not edit:
        return api_error('通过截图为必填项')
    if proof:
        if proof.size > 5 * 1024 * 1024:
            return api_error('截图不能超过 5MB')
        try:
            with warnings.catch_warnings():
                warnings.simplefilter('error', Image.DecompressionBombWarning)
                with Image.open(proof) as picture:
                    if picture.format not in {'PNG', 'JPEG', 'WEBP'}:
                        raise ValueError('unsupported image')
                    if picture.width * picture.height > 20000000:
                        raise ValueError('image too large')
                    extension = {'PNG': '.png', 'JPEG': '.jpg', 'WEBP': '.webp'}[picture.format]
                    picture.verify()
                proof.seek(0)
                with Image.open(proof) as picture:
                    picture.load()
            proof.seek(0)
            import uuid
            proof.name = uuid.uuid4().hex + extension
        except (UnidentifiedImageError, OSError, ValueError, Image.DecompressionBombError, Image.DecompressionBombWarning):
            return api_error('请上传有效的 PNG、JPG 或 WebP 图片')
    logic = request.POST.get('logic', '').strip()
    blog = request.POST.get('blog', '').strip()
    if len(logic) > 50000 or len(blog) > 100000:
        return api_error('逻辑最多 50000 字，Blog 最多 100000 字')
    stored_file = None
    try:
        with transaction.atomic():
            target = Week.objects.filter(pk=week_id, semester__team=membership.team).first()
            if not target:
                return api_error('该周次不存在或不属于当前团队', 404)
            # Serialize submissions across this user's teams and semesters.
            User.objects.select_for_update().get(pk=request.user.pk)
            semester = Semester.objects.select_for_update().get(pk=target.semester_id)
            week = Week.objects.select_for_update().get(pk=target.pk)
            now = timezone.now()
            today = timezone.localdate(now)
            try:
                record_date = datetime.strptime(request.POST.get('record_date') or today.isoformat(), '%Y-%m-%d').date()
            except ValueError:
                return api_error('提交日期格式无效')
            historical = submission_id is None and record_date < today
            if (semester.archived_at or not semester.is_active
                    or not semester.starts_on <= today <= semester.ends_on):
                return api_error('当前学期已结束或尚未开放，不能继续提交', 409)
            if not historical and (week.is_closed or not week.starts_at <= now < week.ends_at):
                return api_error('该周次已截止或尚未开放，不能继续提交，请刷新页面', 409)
            if (week.starts_at, week.ends_at) != week_bounds(semester, week.number):
                return api_error('周次日期异常，请联系管理员修复后再提交', 409)
            if not (timezone.localdate(week.starts_at) <= record_date < timezone.localdate(week.ends_at)
                    and semester.starts_on <= record_date <= min(today, semester.ends_on)):
                return api_error('提交日期不属于所选周次或尚未开放', 409)
            from .models import MakeupPermission, CheckinRecord
            permission = None
            if historical:
                locked_member = TeamMembership.objects.select_for_update().get(pk=membership.pk)
                if not locked_member.is_active:
                    return api_error('当前团队成员权限已失效', 403)
                permission = MakeupPermission.objects.select_for_update().filter(member=membership,
                    semester=semester, record_date=record_date).first()
                if not permission:
                    return api_error('请先消耗补签券开放该日期的提交权限', 409)
                if (permission.completed_at or CheckinRecord.objects.filter(member=membership, checkin_date=record_date).exists()
                        or Submission.objects.filter(member=membership, record_date=record_date, is_valid=True).exists()):
                    return api_error('该日期已完成签到或补交，不能重复提交', 409)
            problem = get_problem(request.POST.get('problem_url', ''))
            duplicates = Submission.objects.filter(member__user_id=request.user.pk, problem=problem)
            if submission_id is not None:
                duplicates = duplicates.exclude(pk=submission_id)
            if duplicates.exists():
                return api_error('你已提交过这道题目，请编辑或重新提交原记录；如需新建，请先删除原记录', 409)
            if submission_id is not None:
                submission = Submission.objects.select_for_update().filter(pk=submission_id, team=membership.team, member=membership, week=week).first()
                if not submission or (not edit and submission.is_valid):
                    return api_error('仅能重交本人原周已退回记录', 409)
                if edit and request.POST.get('version') != str(submission.version):
                    return api_error('记录已变化，请刷新后重新编辑', 409)
                replaced = submission.parts.all() if proof else submission.parts.exclude(kind=SubmissionPart.PROOF)
                if not proof and not submission.parts.filter(kind=SubmissionPart.PROOF, upload__gt='').exists():
                    return api_error('通过截图为必填项')
                FileDeletion.objects.bulk_create([FileDeletion(name=p.upload.name) for p in replaced if p.upload])
                replaced.delete()
                submission.likes.all().delete()
                submission.version += 1
                submission.is_valid = True
                submission.rejection_note = ''
                submission.problem = problem
                # Resubmissions keep the original attributed day.
                submission.save()
                ReviewLog.objects.create(submission=submission, actor=request.user, action='edit' if edit else 'resubmit', detail={'version': submission.version})
            else:
                submission = Submission.objects.create(team=membership.team, member=membership, semester=semester, week=week, is_valid=True, problem=problem, record_date=record_date)
            Submission.objects.filter(pk=submission.pk).update(submitted_at=now)
            submission.submitted_at = now
            if proof:
                part = SubmissionPart.objects.create(submission=submission, kind=SubmissionPart.PROOF, upload=proof)
                stored_file = part.upload
            if logic:
                SubmissionPart.objects.create(submission=submission, kind=SubmissionPart.LOGIC, text_content=logic)
            if blog:
                SubmissionPart.objects.create(submission=submission, kind=SubmissionPart.BLOG, text_content=blog)
            for part in submission.parts.all():
                part.status = SubmissionPart.APPROVED
                part.points_awarded = score_part(part)
                part.save(update_fields=['status', 'points_awarded'])
            if permission:
                CheckinRecord.objects.create(team=membership.team, member=membership,
                    checkin_date=record_date, is_makeup=True)
                permission.completed_at = now
                permission.save(update_fields=['completed_at'])
            refresh_totals(semester)
            if membership.team.currency_enabled:
                from .currency_services import settle_weekly_streak,perform_checkin
                TeamMembership.objects.select_for_update().get(pk=membership.pk)
                if submission.record_date == today and not CheckinRecord.objects.filter(member=membership,checkin_date=today).exists():
                    perform_checkin(membership)
                else:
                    settle_weekly_streak(membership, submission.record_date)
    except Exception as exc:
        if stored_file:
            stored_file.delete(save=False)
        if isinstance(exc,ValueError):return api_error(str(exc),400)
        raise
    return JsonResponse({'ok': True, 'submission': serialize_submission(submission)}, status=201)


@api_login_required
def leaderboard(request):
    membership = membership_for(request)
    if not membership:
        return api_error("未加入当前团队", 403)
    if request.method != 'GET':
        return api_error('仅支持 GET', 405)
    from .rankings import ranking_data
    scope = request.GET.get('scope', 'checkin')
    if scope not in {'checkin', 'rating', 'solved', 'wealth'}:
        return api_error('请选择签到榜、Rating榜、过题榜或财富榜')
    try:
        return JsonResponse(ranking_data(membership, scope, request.GET.get('platform','codeforces')))
    except ValueError as exc:
        return api_error(str(exc))



@api_login_required
@transaction.atomic
def team_settings(request):
    membership = membership_for(request)
    if not membership or not request.user.is_superuser:
        return api_error("需要管理员权限", 403)
    team = membership.team
    if request.method == "GET":
        semester = active_semester(team)
        from .management_views import cleanup_schedule_data
        return JsonResponse({"cleanup": cleanup_schedule_data(semester), "semester": {"minSubmissions": semester.min_submissions, "minScore": semester.min_score, "requireBoth": semester.require_both, "screenshotPoints": semester.screenshot_points, "logicPoints": semester.logic_points, "blogPoints": semester.blog_points} if semester else None})
    if request.method != "PUT":
        return api_error("仅支持 GET 或 PUT", 405)
    data = body(request)
    if 'cleanupDelayDays' in data:
        return api_error('清理天数配置已停用，请在团队配置中设置具体清理时间')
    for key in ['minSubmissions', 'minScore', 'screenshotPoints', 'logicPoints', 'blogPoints']:
        if key in data:
            value = data[key]
            if type(value) is not int or not 0 <= value <= 100000:
                return api_error('设置必须是有效范围内的整数')
    if 'requireBoth' in data and type(data['requireBoth']) is not bool:
        return api_error('达标规则必须为布尔值')
    semester = active_semester(team)
    if semester:
        semester = Semester.objects.select_for_update().get(pk=semester.pk)
        if semester.archived_at:
            return api_error('学期已归档', 409)
        for field, key in (("min_submissions", "minSubmissions"), ("min_score", "minScore"), ("require_both", "requireBoth"), ("screenshot_points", "screenshotPoints"), ("logic_points", "logicPoints"), ("blog_points", "blogPoints")):
            if key in data:
                setattr(semester, field, data[key])
        semester.save()
        refresh_totals(semester)
    return JsonResponse({"ok": True})


@api_login_required
def export_xlsx(request):
    membership = membership_for(request)
    if not membership or not membership.is_admin:
        return api_error("需要管理员权限", 403)
    team = membership.team
    semester_id = request.GET.get('semester')
    if semester_id and not semester_id.isdigit():
        return api_error('学期编号无效')
    semester = team.semesters.filter(id=semester_id).first() if semester_id else team.semesters.order_by('-starts_on').first()
    if not semester:
        return api_error("没有可导出的学期", 404)
    workbook = Workbook()
    sheet = workbook.active
    if semester.archived_at:
        sheet.title = '学期汇总'
        sheet.append(['成员姓名', '学期累计积分', '有效提交次数', '达标周数'])
        for total in semester.member_totals.select_related('member').order_by('member_id'):
            sheet.append([total.display_name or total.member.display_name, total.total_score, total.total_submissions, total.qualified_weeks])
            sheet.cell(sheet.max_row, 1).data_type = 's'
        return FileResponse(_workbook_file(workbook), as_attachment=True, filename=f'{team.code}-archive.xlsx')
    sheet.title = "提交与审核"
    headers = ["成员姓名", "周次", "归属日期", "提交时间", "有效提交", "截图审核", "截图积分", "写题逻辑审核", "逻辑积分", "Blog审核", "Blog积分"]
    sheet.append(headers)
    for submission in team.submissions.filter(semester=semester).select_related("member", "week").prefetch_related("parts"):
        by_kind = {part.kind: part for part in submission.parts.all()}
        sheet.append([submission.member.display_name, submission.week.number, submission.record_date.isoformat(), timezone.localtime(submission.submitted_at).strftime("%Y-%m-%d %H:%M"), "是" if submission.is_valid else "否", by_kind.get("proof").status if by_kind.get("proof") else "", by_kind.get("proof").points_awarded if by_kind.get("proof") else 0, by_kind.get("logic").status if by_kind.get("logic") else "", by_kind.get("logic").points_awarded if by_kind.get("logic") else 0, by_kind.get("blog").status if by_kind.get("blog") else "", by_kind.get("blog").points_awarded if by_kind.get("blog") else 0])
    for row in sheet.iter_rows(min_row=2):
        row[0].data_type = 's'
    for column, value in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(column)].width = max(12, len(value) + 2)
    response = FileResponse(_workbook_file(workbook), as_attachment=True, filename=f"{team.code}-{semester.name}-export.xlsx")
    return response


def _workbook_file(workbook):
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


@api_login_required
def attachment(request, part_id):
    if request.method not in {'GET', 'HEAD'}:
        return api_error('仅支持 GET 或 HEAD', 405)
    part = SubmissionPart.objects.select_related('submission__team', 'submission__semester').filter(pk=part_id).first()
    if not part or not part.upload or part.submission.semester.archived_at:
        raise Http404
    submission = part.submission
    viewer = membership_for(request, submission.team_id)
    if not viewer or not (viewer.is_admin or viewer.id == submission.member_id or
            (submission.is_valid and part.status == SubmissionPart.APPROVED)):
        raise Http404
    if settings.USE_X_ACCEL_REDIRECT:
        response = HttpResponse()
        response['X-Accel-Redirect'] = '/protected-media/' + quote(part.upload.name, safe='/')
    else:
        try:
            response = FileResponse(part.upload.open('rb'))
        except FileNotFoundError:
            raise Http404
    import mimetypes
    content_type = mimetypes.guess_type(part.upload.name)[0]
    if content_type not in {'image/png', 'image/jpeg', 'image/webp'}:
        content_type = 'application/octet-stream'
        response['Content-Disposition'] = 'attachment'
    response['Content-Type'] = content_type
    response['Cache-Control'] = 'private, no-store, max-age=0'
    response['X-Content-Type-Options'] = 'nosniff'
    return response


@api_login_required
def semesters(request):
    membership = membership_for(request)
    if not membership:
        return api_error('未加入当前团队', 403)
    return JsonResponse({'semesters': [{'id': item.id, 'name': item.name, 'archived': bool(item.archived_at),
        'weeks': [week.number for week in item.weeks.all()]}
        for item in membership.team.semesters.prefetch_related('weeks')]})


@api_login_required
@transaction.atomic
def delete_own_submission(request, submission_id):
    if request.method != 'DELETE':
        return api_error('仅支持 DELETE', 405)
    member = membership_for(request)
    target = Submission.objects.filter(pk=submission_id, member=member, team_id=member.team_id).first() if member and member.pk else None
    if not target:
        return api_error('提交不存在或不属于本人', 404)
    semester = Semester.objects.select_for_update().get(pk=target.semester_id)
    Week.objects.select_for_update().get(pk=target.week_id)
    target = Submission.objects.select_for_update().filter(pk=submission_id, member=member).first()
    if not target:
        return api_error('记录已删除，请刷新', 404)
    if semester.archived_at:
        return api_error('学期已归档，不能删除', 409)
    if body(request).get('version') != target.version:
        return api_error('记录已变化，请刷新后重试', 409)
    FileDeletion.objects.bulk_create([FileDeletion(name=part.upload.name) for part in target.parts.all() if part.upload])
    target.delete()
    refresh_totals(semester)
    return JsonResponse({'ok': True})
